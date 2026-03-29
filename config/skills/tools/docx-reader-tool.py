"""
title: Word Document Reader
description: Read and extract text from .docx Word documents. Trigger with: "read this Word file: /path/to/file.docx"
author: LocalAI
version: 1.0
license: MIT
"""

import os


class Tools:

    def read_docx(self, file_path: str) -> str:
        """
        Extracts all text from a Microsoft Word (.docx) document.
        Trigger phrases: read this Word file, open docx, summarize Word document, read .docx
        :param file_path: Path to the .docx file (e.g. C:/Users/You/Documents/report.docx)
        :return: Extracted text from the document
        """
        try:
            from docx import Document
        except ImportError:
            return (
                "❌ **python-docx not installed.**\n\n"
                "Run `install-skill-packages.bat` in the Local AI folder to install it.\n"
                "Then try again."
            )

        if not os.path.isfile(file_path):
            return (
                f"❌ File not found: `{file_path}`\n\n"
                "Make sure the path is correct.\n"
                "Example: `C:/Users/You/Documents/report.docx`"
            )

        ext = os.path.splitext(file_path)[1].lower()
        if ext != ".docx":
            return f"❌ This tool only reads .docx files. Got: `{ext}`\nFor .doc files, please convert to .docx first."

        try:
            doc = Document(file_path)

            sections = []

            # Extract heading and paragraph structure
            for para in doc.paragraphs:
                if para.style.name.startswith("Heading"):
                    level = para.style.name.replace("Heading ", "")
                    prefix = "#" * min(int(level), 6) if level.isdigit() else "##"
                    sections.append(f"\n{prefix} {para.text.strip()}")
                elif para.text.strip():
                    sections.append(para.text.strip())

            # Extract tables
            table_count = 0
            for table in doc.tables:
                table_count += 1
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    sections.append(f"\n**Table {table_count}:**")
                    sections.extend(rows)

            if not sections:
                return f"⚠️ No text found in `{os.path.basename(file_path)}`."

            full_text = "\n".join(sections).strip()
            word_count = len(full_text.split())

            return (
                f"✅ **Document extracted:** `{os.path.basename(file_path)}`\n"
                f"Words: {word_count} | Tables: {table_count}\n\n"
                f"---\n\n{full_text}"
            )

        except Exception as e:
            return f"❌ Error reading document: {e}"
