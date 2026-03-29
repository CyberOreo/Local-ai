"""
title: Excel Spreadsheet Reader
description: Read and analyze Excel .xlsx files. Trigger with: "analyze this spreadsheet: /path/to/file.xlsx"
author: LocalAI
version: 1.0
license: MIT
"""

import os


class Tools:

    def read_xlsx(self, file_path: str, max_rows: int = 100) -> str:
        """
        Reads an Excel spreadsheet and returns its contents with basic statistics.
        Trigger phrases: analyze this spreadsheet, read Excel file, open xlsx, summarize this Excel
        :param file_path: Path to the .xlsx file
        :param max_rows: Max rows to display per sheet (default 100)
        :return: Spreadsheet contents and basic stats
        """
        try:
            import openpyxl
        except ImportError:
            return (
                "❌ **openpyxl not installed.**\n\n"
                "Run `install-skill-packages.bat` in the Local AI folder to install it.\n"
                "Then try again."
            )

        if not os.path.isfile(file_path):
            return f"❌ File not found: `{file_path}`"

        ext = os.path.splitext(file_path)[1].lower()
        if ext != ".xlsx":
            return f"❌ This tool reads .xlsx files only. Got: `{ext}`"

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            output = [
                f"✅ **Spreadsheet:** `{os.path.basename(file_path)}`\n"
                f"Sheets: {len(wb.sheetnames)} — {', '.join(wb.sheetnames)}\n"
            ]

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))

                if not rows:
                    output.append(f"\n### Sheet: {sheet_name}\n*Empty sheet*")
                    continue

                total_rows = len(rows)
                total_cols = max(len(r) for r in rows) if rows else 0

                output.append(f"\n### Sheet: {sheet_name}")
                output.append(f"*{total_rows} rows × {total_cols} columns*\n")

                display_rows = rows[:max_rows]

                # Format as markdown table
                table_lines = []
                for i, row in enumerate(display_rows):
                    cells = [str(c) if c is not None else "" for c in row]
                    table_lines.append(" | ".join(cells))
                    if i == 0:
                        # Add header separator
                        table_lines.append(" | ".join(["---"] * len(cells)))

                output.extend(table_lines)

                if total_rows > max_rows:
                    output.append(f"\n*...{total_rows - max_rows} more rows not shown*")

                # Basic numeric stats
                numeric_cols = {}
                for col_idx in range(total_cols):
                    values = []
                    for row in rows[1:]:  # skip header
                        if col_idx < len(row) and isinstance(row[col_idx], (int, float)):
                            values.append(row[col_idx])
                    if values:
                        header_val = rows[0][col_idx] if col_idx < len(rows[0]) else f"Col {col_idx+1}"
                        numeric_cols[str(header_val)] = {
                            "count": len(values),
                            "sum": sum(values),
                            "min": min(values),
                            "max": max(values),
                            "avg": round(sum(values) / len(values), 2)
                        }

                if numeric_cols:
                    output.append("\n**Numeric column stats:**")
                    for col_name, stats in list(numeric_cols.items())[:8]:
                        output.append(
                            f"- **{col_name}**: count={stats['count']}, "
                            f"sum={stats['sum']}, min={stats['min']}, "
                            f"max={stats['max']}, avg={stats['avg']}"
                        )

            wb.close()
            return "\n".join(output)

        except Exception as e:
            return f"❌ Error reading spreadsheet: {e}"
